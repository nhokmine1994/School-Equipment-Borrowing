#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script sinh file PowerPoint (PPTX) chất lượng cao với 100% NATIVE EDITABLE OBJECTS.
Tất cả các khối văn bản, thẻ bo góc, số liệu thống kê, thanh quy trình đều là đối tượng có thể chỉnh sửa trực tiếp trong Microsoft PowerPoint.
Tông màu: Chuẩn giao diện trường học SEB (Primary Blue #1976D2, Navy #0B5AA9, Sky #42A5F5, Gold #E0CE92, White #FFFFFF).
"""

import os
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

# Palette chuẩn của dự án School-Equipment-Borrowing
COLOR_PRIMARY_BLUE = RGBColor(25, 118, 210)   # #1976D2
COLOR_NAVY_DARK    = RGBColor(11, 90, 169)    # #0B5AA9
COLOR_SKY_BLUE     = RGBColor(66, 165, 245)   # #42A5F5
COLOR_LIGHT_BG     = RGBColor(240, 247, 252)  # #F0F7FC
COLOR_CONTAINER_BG = RGBColor(229, 243, 253)  # #E5F3FD
COLOR_GOLD_ACCENT  = RGBColor(224, 206, 146)  # #E0CE92
COLOR_WHITE        = RGBColor(255, 255, 255)  # #FFFFFF
COLOR_TEXT_DARK    = RGBColor(15, 23, 42)     # #0F172A
COLOR_TEXT_MUTED   = RGBColor(71, 85, 105)    # #475569
COLOR_BORDER_LIGHT = RGBColor(186, 224, 253)  # #BAE0FD
COLOR_SUCCESS      = RGBColor(16, 185, 129)   # #10B981
COLOR_AMBER        = RGBColor(245, 158, 11)   # #F59E0B
COLOR_RED          = RGBColor(239, 68, 68)    # #EF4444

def set_shape_flat(shape, fill_color, border_color=None, border_width=Pt(1)):
    """Đặt màu nền phẳng và viền cho Shape."""
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = border_width
    else:
        shape.line.fill.background()

def add_header(slide, title_text, category_text=None):
    """Tạo Header chuẩn phong cách SEB cho từng slide."""
    # Top banner nhỏ trang trí
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    set_shape_flat(top_bar, COLOR_PRIMARY_BLUE)

    # Khung text tiêu đề
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.9))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    if category_text:
        p_cat = tf.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.name = "Segoe UI"
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_SKY_BLUE
        p_title = tf.add_paragraph()
    else:
        p_title = tf.paragraphs[0]

    p_title.text = title_text
    p_title.font.name = "Segoe UI"
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_NAVY_DARK

def create_presentation(output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5) # Chuẩn 16:9 widescreen
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: BÌA BÁO CÁO (COVER SLIDE)
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    set_shape_flat(bg1, COLOR_LIGHT_BG)

    # Banner xanh chủ đạo phía trên
    header_box = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(2.8))
    set_shape_flat(header_box, COLOR_PRIMARY_BLUE)

    # Dải viền vàng đặc trưng SEB
    gold_line = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(2.8), Inches(13.333), Inches(0.1))
    set_shape_flat(gold_line, COLOR_GOLD_ACCENT)

    # Tiêu đề trên bìa
    tb_title = slide1.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.3), Inches(1.8))
    tf_title = tb_title.text_frame
    tf_title.word_wrap = True

    p0 = tf_title.paragraphs[0]
    p0.text = "BÁO CÁO KỸ THUẬT & ĐỒ ÁN TỐT NGHIỆP"
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_GOLD_ACCENT

    p1 = tf_title.add_paragraph()
    p1.text = "HỆ THỐNG QUẢN LÝ MƯỢN TRẢ THIẾT BỊ (SEB)"
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_WHITE

    p2 = tf_title.add_paragraph()
    p2.text = "Đơn vị thụ hưởng: Trường THCS Lộc An | Mô hình Hybrid React - PHP & SQL Server"
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(15)
    p2.font.color.rgb = COLOR_WHITE

    # Khung thẻ thành viên thực hiện
    card_info = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(3.4), Inches(11.333), Inches(3.4))
    set_shape_flat(card_info, COLOR_WHITE, COLOR_BORDER_LIGHT, Pt(1.5))

    tb_team = slide1.shapes.add_textbox(Inches(1.5), Inches(3.7), Inches(10.3), Inches(2.8))
    tf_team = tb_team.text_frame
    tf_team.word_wrap = True

    pt0 = tf_team.paragraphs[0]
    pt0.text = "THÔNG TIN THÀNH VIÊN NHÓM THỰC HIỆN:"
    pt0.font.name = "Segoe UI"
    pt0.font.size = Pt(16)
    pt0.font.bold = True
    pt0.font.color.rgb = COLOR_NAVY_DARK

    members = [
        ("Phạm Công Hậu", "MSSV: 525000458", "Trưởng nhóm / Kiến trúc hệ thống & Backend PHP"),
        ("Lò Văn Duẩn", "MSSV: 525000631", "Lập trình Giao diện React SPA & Tối ưu UI/UX"),
        ("Đặng Bắc Nam", "MSSV: 425000218", "Thiết kế CSDL SQL Server, Nghiệm thu & Tích hợp")
    ]

    for name, mssv, role in members:
        pm = tf_team.add_paragraph()
        pm.space_before = Pt(8)
        r_name = pm.add_run()
        r_name.text = f"• {name} "
        r_name.font.name = "Segoe UI"
        r_name.font.size = Pt(14)
        r_name.font.bold = True
        r_name.font.color.rgb = COLOR_PRIMARY_BLUE

        r_mssv = pm.add_run()
        r_mssv.text = f"({mssv}) - "
        r_mssv.font.name = "Segoe UI"
        r_mssv.font.size = Pt(14)
        r_mssv.font.bold = True
        r_mssv.font.color.rgb = COLOR_TEXT_DARK

        r_role = pm.add_run()
        r_role.text = role
        r_role.font.name = "Segoe UI"
        r_role.font.size = Pt(13)
        r_role.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 2: INFOGRAPHIC TỔNG QUAN HỆ THỐNG (100% NATIVE EDITABLE OBJECTS)
    # Đây chính là slide thiết kế lại từ hình của người dùng!
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    bg2 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    set_shape_flat(bg2, COLOR_LIGHT_BG)

    # Tiêu đề Infographic
    add_header(slide2, "Hệ Thống Quản Lý Mượn Trả Thiết Bị (SEB) - Trường THCS Lộc An", "INFOGRAPHIC TỔNG QUAN GIẢI PHÁP SỐ HÓA")

    # Phụ đề mô tả
    tb_sub = slide2.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(0.45))
    tf_sub = tb_sub.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "Giải pháp số hóa toàn diện quản lý kho thiết bị, phòng học và tin tức tại Trường THCS Lộc An. Vận hành trên nền tảng Hybrid React-PHP hiện đại, bảo mật và hiệu suất cao."
    p_sub.font.name = "Segoe UI"
    p_sub.font.size = Pt(12)
    p_sub.font.color.rgb = COLOR_TEXT_MUTED

    # CỘT TRÁI: TÍNH NĂNG DÀNH CHO NGƯỜI DÙNG (Width: 5.7 inches)
    col1_left = Inches(0.8)
    col1_width = Inches(5.7)

    # Header Cột Trái
    card_hdr1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col1_left, Inches(1.85), col1_width, Inches(0.45))
    set_shape_flat(card_hdr1, COLOR_PRIMARY_BLUE)
    p_hdr1 = card_hdr1.text_frame.paragraphs[0]
    p_hdr1.text = "TÍNH NĂNG DÀNH CHO NGƯỜI DÙNG"
    p_hdr1.font.name = "Segoe UI"
    p_hdr1.font.size = Pt(12)
    p_hdr1.font.bold = True
    p_hdr1.font.color.rgb = COLOR_WHITE
    p_hdr1.alignment = PP_ALIGN.CENTER

    # Thẻ 1.1: Quản lý thiết bị thời gian thực
    c1_1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col1_left, Inches(2.4), col1_width, Inches(1.5))
    set_shape_flat(c1_1, COLOR_WHITE, COLOR_BORDER_LIGHT, Pt(1))
    tf1_1 = c1_1.text_frame
    tf1_1.word_wrap = True
    tf1_1.margin_left = tf1_1.margin_top = tf1_1.margin_right = Inches(0.2)
    p = tf1_1.paragraphs[0]
    p.text = "Quản lý thiết bị thời gian thực"
    p.font.name = "Segoe UI"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY_DARK

    p = tf1_1.add_paragraph()
    p.text = "Theo dõi trạng thái của hơn 600 thiết bị từ máy tính đến dụng cụ thí nghiệm."
    p.font.name = "Segoe UI"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_TEXT_MUTED

    # 3 Khối Badge số liệu thời gian thực (Native editable pill badges)
    stat_boxes = [
        ("631", "Hiện có", COLOR_PRIMARY_BLUE),
        ("65", "Đang mượn", COLOR_AMBER),
        ("0", "Bảo trì", COLOR_SUCCESS)
    ]
    b_w = Inches(1.65)
    for idx, (num, label, col) in enumerate(stat_boxes):
        bx = col1_left + Inches(0.25) + idx * (b_w + Inches(0.15))
        badge = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, Inches(3.2), b_w, Inches(0.55))
        set_shape_flat(badge, COLOR_CONTAINER_BG, col, Pt(1.5))
        tb_bg = badge.text_frame
        tb_bg.margin_top = Inches(0.04)
        p_b1 = tb_bg.paragraphs[0]
        p_b1.text = f"{num} {label}"
        p_b1.font.name = "Segoe UI"
        p_b1.font.size = Pt(11)
        p_b1.font.bold = True
        p_b1.font.color.rgb = col
        p_b1.alignment = PP_ALIGN.CENTER

    # Thẻ 1.2: Đăng ký phòng học trực quan
    c1_2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col1_left, Inches(4.0), col1_width, Inches(1.4))
    set_shape_flat(c1_2, COLOR_WHITE, COLOR_BORDER_LIGHT, Pt(1))
    tf1_2 = c1_2.text_frame
    tf1_2.word_wrap = True
    tf1_2.margin_left = tf1_2.margin_top = tf1_2.margin_right = Inches(0.2)
    p = tf1_2.paragraphs[0]
    p.text = "Đăng ký phòng học trực quan"
    p.font.name = "Segoe UI"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY_DARK

    p = tf1_2.add_paragraph()
    p.text = "Lịch biểu phân chia theo tiết học và ca học (Sáng / Chiều T2–T6) giúp giáo viên dễ dàng kiểm tra phòng trống và đặt lịch giảng dạy tức thì."
    p.font.name = "Segoe UI"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_TEXT_MUTED

    # Thẻ 1.3: Cá nhân hóa kho lưu trữ
    c1_3 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col1_left, Inches(5.5), col1_width, Inches(1.5))
    set_shape_flat(c1_3, COLOR_WHITE, COLOR_BORDER_LIGHT, Pt(1))
    tf1_3 = c1_3.text_frame
    tf1_3.word_wrap = True
    tf1_3.margin_left = tf1_3.margin_top = tf1_3.margin_right = Inches(0.2)
    p = tf1_3.paragraphs[0]
    p.text = "Cá nhân hóa kho lưu trữ (Kho Cá Nhân)"
    p.font.name = "Segoe UI"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY_DARK

    p = tf1_3.add_paragraph()
    p.text = "Mỗi người dùng có không gian riêng để theo dõi lịch sử mượn, thiết bị đang giữ, thời hạn hoàn trả và trạng thái duyệt phiếu của từng đơn hàng."
    p.font.name = "Segoe UI"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_TEXT_MUTED

    # CỘT PHẢI: NỀN TẢNG & QUẢN TRỊ HỆ THỐNG (Width: 5.7 inches)
    col2_left = Inches(6.8)
    col2_width = Inches(5.7)

    # Header Cột Phải
    card_hdr2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col2_left, Inches(1.85), col2_width, Inches(0.45))
    set_shape_flat(card_hdr2, COLOR_NAVY_DARK)
    p_hdr2 = card_hdr2.text_frame.paragraphs[0]
    p_hdr2.text = "NỀN TẢNG & QUẢN TRỊ HỆ THỐNG"
    p_hdr2.font.name = "Segoe UI"
    p_hdr2.font.size = Pt(12)
    p_hdr2.font.bold = True
    p_hdr2.font.color.rgb = COLOR_WHITE
    p_hdr2.alignment = PP_ALIGN.CENTER

    # Thẻ 2.1: Kiến trúc Hybrid React-PHP
    c2_1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col2_left, Inches(2.4), col2_width, Inches(1.5))
    set_shape_flat(c2_1, COLOR_WHITE, COLOR_BORDER_LIGHT, Pt(1))
    tf2_1 = c2_1.text_frame
    tf2_1.word_wrap = True
    tf2_1.margin_left = tf2_1.margin_top = tf2_1.margin_right = Inches(0.2)
    p = tf2_1.paragraphs[0]
    p.text = "Kiến trúc Hybrid React-PHP"
    p.font.name = "Segoe UI"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY_DARK

    p = tf2_1.add_paragraph()
    p.text = "Sử dụng React cho giao diện mượt mà (SPA, Vite build, Skeleton loading) kết hợp PHP phía Backend để xử lý xác thực session an toàn và kết nối cơ sở dữ liệu."
    p.font.name = "Segoe UI"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_TEXT_MUTED

    # Thẻ 2.2: Quy trình duyệt 3 bước
    c2_2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col2_left, Inches(4.0), col2_width, Inches(1.4))
    set_shape_flat(c2_2, COLOR_WHITE, COLOR_BORDER_LIGHT, Pt(1))
    tf2_2 = c2_2.text_frame
    tf2_2.word_wrap = True
    tf2_2.margin_left = tf2_2.margin_top = tf2_2.margin_right = Inches(0.2)
    p = tf2_2.paragraphs[0]
    p.text = "Quy trình duyệt 3 bước chuẩn hóa"
    p.font.name = "Segoe UI"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY_DARK

    p = tf2_2.add_paragraph()
    p.text = "Mọi yêu cầu được phân loại rõ ràng trong CSDL SQL Server với bảng TinhTrangDuyet:"
    p.font.name = "Segoe UI"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_TEXT_MUTED

    # 3 Khối Mũi tên Chevron / Badge quy trình
    steps = [
        ("1. Chờ duyệt", COLOR_AMBER),
        ("2. Đã duyệt", COLOR_SUCCESS),
        ("3. Từ chối", COLOR_RED)
    ]
    st_w = Inches(1.65)
    for idx, (st_text, col) in enumerate(steps):
        sx = col2_left + Inches(0.25) + idx * (st_w + Inches(0.15))
        badge = slide2.shapes.add_shape(MSO_SHAPE.CHEVRON, sx, Inches(4.8), st_w, Inches(0.45))
        set_shape_flat(badge, col)
        p_st = badge.text_frame.paragraphs[0]
        p_st.text = st_text
        p_st.font.name = "Segoe UI"
        p_st.font.size = Pt(11)
        p_st.font.bold = True
        p_st.font.color.rgb = COLOR_WHITE
        p_st.alignment = PP_ALIGN.CENTER

    # Thẻ 2.3: Hệ quản trị cơ sở dữ liệu chuyên sâu
    c2_3 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col2_left, Inches(5.5), col2_width, Inches(1.5))
    set_shape_flat(c2_3, COLOR_WHITE, COLOR_BORDER_LIGHT, Pt(1))
    tf2_3 = c2_3.text_frame
    tf2_3.word_wrap = True
    tf2_3.margin_left = tf2_3.margin_top = tf2_3.margin_right = Inches(0.2)
    p = tf2_3.paragraphs[0]
    p.text = "Hệ quản trị cơ sở dữ liệu chuyên sâu (SQL Server)"
    p.font.name = "Segoe UI"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY_DARK

    p = tf2_3.add_paragraph()
    p.text = "Dữ liệu chuẩn hóa trên Microsoft SQL Server với hệ thống Stored Procedure bảo mật, ràng buộc khóa ngoại (Foreign Keys) và quy tắc Default constraint đảm bảo tính toàn vẹn."
    p.font.name = "Segoe UI"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 3: KIẾN TRÚC HYBRID REACT-PHP & BACKEND CHUYÊN SÂU
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    bg3 = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    set_shape_flat(bg3, COLOR_LIGHT_BG)
    add_header(slide3, "Kiến Trúc Công Nghệ Hybrid React + PHP", "KIẾN TRÚC HỆ THỐNG & TECH STACK")

    # 3 Khối Tầng Kiến trúc: Frontend - Backend - Database
    layers = [
        ("TẦNG TRÌNH DIỄN (FRONTEND SPA)", "React, Vite, HTML5, CSS3, Bootstrap", [
            "Giao diện hiện đại, phản hồi tức thì (SPA).",
            "Sử dụng Vite làm công cụ build tối ưu tài nguyên.",
            "Tích hợp Skeleton Loading & Spinner tăng trải nghiệm.",
            "Responsive hoàn chỉnh trên cả Desktop, Tablet & Mobile."
        ], COLOR_PRIMARY_BLUE),
        ("TẦNG XỬ LÝ (BACKEND LOGIC)", "PHP (Native API & Session Helper)", [
            "Đảm nhận định tuyến an toàn qua Apache .htaccess.",
            "Quản lý session người dùng và xác thực phân quyền.",
            "Xử lý logic mượn trả, kiểm tra xung đột thời gian.",
            "Cung cấp fallback an toàn cho môi trường production."
        ], COLOR_NAVY_DARK),
        ("TẦNG DỮ LIỆU (DATABASE TẬP TRUNG)", "Microsoft SQL Server (MSSQL)", [
            "Lưu trữ danh mục 600+ thiết bị, 9 nhóm phân loại.",
            "Bảng TinhTrangDuyet quản lý phân cấp trạng thái duyệt.",
            "Ràng buộc toàn vẹn khóa ngoại & Default constraints.",
            "Sẵn sàng mở rộng hàng đợi thông báo (NotificationQueue)."
        ], COLOR_SKY_BLUE)
    ]

    card_w = Inches(3.64)
    for i, (l_title, l_tech, l_points, l_col) in enumerate(layers):
        cx = Inches(0.8) + i * (card_w + Inches(0.4))
        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, Inches(1.6), card_w, Inches(5.2))
        set_shape_flat(card, COLOR_WHITE, COLOR_BORDER_LIGHT, Pt(1.5))

        # Header card
        ch = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, Inches(1.6), card_w, Inches(0.9))
        set_shape_flat(ch, l_col)
        tf_ch = ch.text_frame
        tf_ch.word_wrap = True
        p1 = tf_ch.paragraphs[0]
        p1.text = l_title
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_WHITE
        p1.alignment = PP_ALIGN.CENTER

        p2 = tf_ch.add_paragraph()
        p2.text = l_tech
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = COLOR_GOLD_ACCENT
        p2.alignment = PP_ALIGN.CENTER

        # Body points
        tb_body = slide3.shapes.add_textbox(cx + Inches(0.2), Inches(2.7), card_w - Inches(0.4), Inches(3.8))
        tf_body = tb_body.text_frame
        tf_body.word_wrap = True
        tf_body.margin_left = tf_body.margin_top = 0

        for pt_idx, pt in enumerate(l_points):
            p_item = tf_body.paragraphs[0] if pt_idx == 0 else tf_body.add_paragraph()
            p_item.space_before = Pt(10)
            p_item.text = f"✔  {pt}"
            p_item.font.name = "Segoe UI"
            p_item.font.size = Pt(11.5)
            p_item.font.color.rgb = COLOR_TEXT_DARK

    # =========================================================================
    # SLIDE 4: QUY TRÌNH NGHIỆP VỤ MƯỢN TRẢ 5 BƯỚC KHÉP KÍN
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    bg4 = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    set_shape_flat(bg4, COLOR_LIGHT_BG)
    add_header(slide4, "Quy Trình Nghiệp Vụ Mượn Trả Thiết Bị 5 Bước Khép Kín", "NGHIỆP VỤ VẬN HÀNH")

    workflow_steps = [
        ("BƯỚC 1", "Tra Cứu & Chọn", "Giáo viên duyệt kho thiết bị theo 9 danh mục hoặc kiểm tra lưới lịch phòng học.", COLOR_PRIMARY_BLUE),
        ("BƯỚC 2", "Tạo Yêu Cầu", "Nhấn nút 'Mượn' và điền mục đích; phiếu mượn tự động lưu với trạng thái Chờ duyệt.", COLOR_SKY_BLUE),
        ("BƯỚC 3", "Xét Duyệt Kho", "Cán bộ kho kiểm tra trên admin_borrows.php và chọn duyệt hoặc từ chối phiếu.", COLOR_AMBER),
        ("BƯỚC 4", "Bàn Giao Đồ", "Giáo viên nhận thiết bị; hệ thống tự động trừ kho và chuyển trạng thái 'Đang mượn'.", COLOR_NAVY_DARK),
        ("BƯỚC 5", "Hoàn Trả & Lưu", "Giáo viên trả thiết bị, kho xác nhận 'Đã trả', hệ thống cộng lại kho và đổi 'Sẵn sàng'.", COLOR_SUCCESS)
    ]

    sw = Inches(2.25)
    for i, (b_name, b_title, b_desc, b_col) in enumerate(workflow_steps):
        bx = Inches(0.8) + i * (sw + Inches(0.12))
        
        # Thẻ bước
        scard = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, Inches(1.8), sw, Inches(4.8))
        set_shape_flat(scard, COLOR_WHITE, COLOR_BORDER_LIGHT, Pt(1.5))

        # Header bước
        sh = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, Inches(1.8), sw, Inches(1.1))
        set_shape_flat(sh, b_col)
        tf_sh = sh.text_frame
        tf_sh.word_wrap = True
        
        p = tf_sh.paragraphs[0]
        p.text = b_name
        p.font.name = "Segoe UI"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE
        p.alignment = PP_ALIGN.CENTER

        p = tf_sh.add_paragraph()
        p.text = b_title
        p.font.name = "Segoe UI"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = COLOR_GOLD_ACCENT
        p.alignment = PP_ALIGN.CENTER

        # Mô tả bước
        s_desc = slide4.shapes.add_textbox(bx + Inches(0.15), Inches(3.2), sw - Inches(0.3), Inches(3.0))
        tf_s = s_desc.text_frame
        tf_s.word_wrap = True
        p_desc = tf_s.paragraphs[0]
        p_desc.text = b_desc
        p_desc.font.name = "Segoe UI"
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = COLOR_TEXT_MUTED
        p_desc.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 5: KẾT QUẢ NGHIỆM THU & ĐÁNH GIÁ CHẤT LƯỢNG
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    bg5 = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    set_shape_flat(bg5, COLOR_LIGHT_BG)
    add_header(slide5, "Báo Cáo Nghiệm Thu & Đánh Giá Chất Lượng Phần Mềm", "KIỂM THỬ VÀ NGHIỆM THU")

    metrics = [
        ("130+", "Tài nguyên Assets liên kết", "100% đường dẫn CSS, JS, hình ảnh được chuẩn hóa.", COLOR_PRIMARY_BLUE),
        ("0%", "Tỷ lệ liên kết hỏng", "Không còn lỗi thiếu ảnh đại diện hay tài nguyên tĩnh.", COLOR_SUCCESS),
        ("631", "Thiết bị được số hóa", "Đầy đủ thông tin môn học, danh mục và tình trạng.", COLOR_SKY_BLUE),
        ("100%", "Chuẩn hóa Responsive", "Hoạt động hoàn hảo trên mọi kích thước màn hình.", COLOR_NAVY_DARK)
    ]

    mw = Inches(2.78)
    for idx, (m_val, m_title, m_desc, m_col) in enumerate(metrics):
        mx = Inches(0.8) + idx * (mw + Inches(0.2))
        m_card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, mx, Inches(1.8), mw, Inches(2.2))
        set_shape_flat(m_card, COLOR_WHITE, COLOR_BORDER_LIGHT, Pt(1.5))

        tf_m = m_card.text_frame
        tf_m.word_wrap = True
        tf_m.margin_left = tf_m.margin_right = tf_m.margin_top = Inches(0.15)
        
        p = tf_m.paragraphs[0]
        p.text = m_val
        p.font.name = "Segoe UI"
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = m_col
        p.alignment = PP_ALIGN.CENTER

        p = tf_m.add_paragraph()
        p.text = m_title
        p.font.name = "Segoe UI"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_DARK
        p.alignment = PP_ALIGN.CENTER

        p = tf_m.add_paragraph()
        p.text = m_desc
        p.font.name = "Segoe UI"
        p.font.size = Pt(9.5)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.alignment = PP_ALIGN.CENTER

    # Bảng kết quả kiểm thử chức năng bên dưới
    summary_box = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.3), Inches(11.733), Inches(2.5))
    set_shape_flat(summary_box, COLOR_WHITE, COLOR_BORDER_LIGHT, Pt(1.5))
    tf_sb = summary_box.text_frame
    tf_sb.word_wrap = True
    tf_sb.margin_left = tf_sb.margin_top = Inches(0.3)

    p = tf_sb.paragraphs[0]
    p.text = "TỔNG HỢP KIỂM THỬ THỰC TẾ TRÊN MÔI TRƯỜNG LOCALHOST / XAMPP:"
    p.font.name = "Segoe UI"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY_DARK

    test_lines = [
        "1. Xác thực & Phân quyền: Đảm bảo phân tách rạch ròi giữa giao diện Giáo viên và Bảng quản trị Admin.",
        "2. Toàn vẹn dữ liệu: Bảng TinhTrangDuyet ngăn chặn triệt để tình trạng đơn mượn bị sai lệch mã trạng thái.",
        "3. Tối ưu hiệu năng: Cơ chế tải dữ liệu dạng chỉ-đọc kết hợp Skeleton Screen triệt tiêu hiện tượng lag trang.",
        "4. Tương thích triển khai: Hệ thống vận hành ổn định trên máy chủ mới thông qua cấu hình biến môi trường kết nối."
    ]
    for line in test_lines:
        p = tf_sb.add_paragraph()
        p.space_before = Pt(4)
        p.text = f"• {line}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT_DARK

    # =========================================================================
    # SLIDE 6: LỘ TRÌNH PHÁT TRIỂN TƯƠNG LAI (ROADMAP)
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    bg6 = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    set_shape_flat(bg6, COLOR_LIGHT_BG)
    add_header(slide6, "Lộ Trình Phát Triển Tương Lai (Roadmap)", "HƯỚNG MỞ RỘNG VÀ NÂNG CẤP")

    roadmap = [
        ("SMART BORROWING (QR)", "Quản lý mã QR thông minh", "Dán nhãn mã QR trên từng thiết bị, giáo viên quét mã bằng điện thoại để hoàn tất thủ tục mượn/trả trong 3 giây.", COLOR_PRIMARY_BLUE),
        ("AI ASSISTANT", "Trợ lý ảo hỏi đáp", "Tích hợp trợ lý AI thông minh giải đáp quy định, gợi ý thiết bị thay thế và tự động lập báo cáo kiểm kê kho định kỳ.", COLOR_SKY_BLUE),
        ("OMNI NOTIFICATIONS", "Hệ thống thông báo đa kênh", "Tự động gửi thông báo duyệt phiếu, nhắc lịch trả thiết bị qua Zalo ZNS, Email và thông báo trên trình duyệt.", COLOR_AMBER),
        ("MULTI-SCHOOL PLATFORM", "Nền tảng liên trường", "Mở rộng hệ thống để quản lý liên thông giữa nhiều trường học trên địa bàn, chia sẻ kho thiết bị dùng chung.", COLOR_SUCCESS)
    ]

    rw = Inches(2.78)
    for idx, (r_code, r_title, r_desc, r_col) in enumerate(roadmap):
        rx = Inches(0.8) + idx * (rw + Inches(0.2))
        rcard = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, rx, Inches(1.8), rw, Inches(5.0))
        set_shape_flat(rcard, COLOR_WHITE, COLOR_BORDER_LIGHT, Pt(1.5))

        # Header Roadmap
        rh = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, rx, Inches(1.8), rw, Inches(1.1))
        set_shape_flat(rh, r_col)
        tf_rh = rh.text_frame
        tf_rh.word_wrap = True
        
        p = tf_rh.paragraphs[0]
        p.text = r_code
        p.font.name = "Segoe UI"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_GOLD_ACCENT
        p.alignment = PP_ALIGN.CENTER

        p = tf_rh.add_paragraph()
        p.text = r_title
        p.font.name = "Segoe UI"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE
        p.alignment = PP_ALIGN.CENTER

        # Body
        rb = slide6.shapes.add_textbox(rx + Inches(0.15), Inches(3.1), rw - Inches(0.3), Inches(3.4))
        tf_rb = rb.text_frame
        tf_rb.word_wrap = True
        p_desc = tf_rb.paragraphs[0]
        p_desc.text = r_desc
        p_desc.font.name = "Segoe UI"
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = COLOR_TEXT_MUTED
        p_desc.alignment = PP_ALIGN.CENTER

    # Lưu file PPTX
    prs.save(output_path)
    print(f"[✓] Đã tạo thành công file PowerPoint Editable tại: {output_path}")

if __name__ == "__main__":
    output_dir = Path(r"d:\School-Equipment-Borrowing\.notebooklm\downloads\Hệ Thống Mượn Trả Thiết Bị Trường Lộc An")
    output_dir.mkdir(parents=True, exist_ok=True)
    out_file = output_dir / "SEB_Bao_Cao_Thuyet_Trinh_EDITABLE.pptx"
    create_presentation(str(out_file))
